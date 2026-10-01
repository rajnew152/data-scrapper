"""Multi-user concurrency tests (offline, no credits, no network).

Part 1 - JobManager rules with instant fake jobs:
    slot limit + FIFO, one job per category, later jobs for free categories
    overtake a blocked one, per-user limit, queue cap, stop while queued,
    stop while running, bounded history, a dead job thread never holds a slot
Part 2 - the REAL engine, several jobs at once against simulated Serper +
    websites: every job completes, no business is stored twice across the
    concurrent jobs, the account-wide Serper in-flight cap holds across
    jobs, the coalesced checkpoint and the Excel files hold every record.

Run (from app/):  python tests/test_concurrency.py
"""

from __future__ import annotations

import os
import sys
import tempfile
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, os.path.join(HERE, "loadtest"))

from backend import config  # noqa: E402

_tmp = tempfile.mkdtemp(prefix="concurrency_test_")
config.OUTPUT_DIR = _tmp
config.STATE_PATH = os.path.join(_tmp, "demo_state.json")
config.STATE_SAVE_INTERVAL = 0.5
config.EXPORT_MIN_INTERVAL = 2.0
config.MAX_CREDITS_PER_RECORD = 0
config.SERPER_RPS = 50
config.SERPER_CONCURRENCY = 6

import sim  # noqa: E402

from backend.collector.engine import StateStore  # noqa: E402
from backend.collector.jobs import JobManager, JobRejected  # noqa: E402
from backend.collector.normalize import norm_domain, norm_name  # noqa: E402


# --------------------------------------------------------------------------- #
# Part 1: scheduling rules
# --------------------------------------------------------------------------- #

class FakeJob:
    """The slice of CollectionJob the manager uses; runs until released."""

    def __init__(self, state, category, keywords, location, target, provider,
                 max_queries, geo=None):
        import uuid
        self.id = uuid.uuid4().hex[:16]
        self.category, self.target, self.display = category, target, category
        self.owner = ""
        self.status = "pending"
        self.stop_reason = self.error = self.finished_at = ""
        self.created_at = time.strftime("%Y-%m-%dT%H:%M:%S")
        self.stop_event = threading.Event()
        self.release = threading.Event()
        self.thread = None
        self.on_finish = None
        self.crawl_workers = 0
        self.search_share = lambda: 1
        self.lines: list[str] = []

    def say(self, msg):
        self.lines.append(msg)

    def category_count(self):
        return 0

    def start(self):
        self.thread = threading.Thread(target=self._run, daemon=True)
        self.thread.start()

    def _run(self):
        self.status = "running"
        while not (self.release.is_set() or self.stop_event.is_set()):
            time.sleep(0.01)
        self.status = "stopped" if self.stop_event.is_set() else "completed"
        if self.on_finish:
            self.on_finish(self)

    def stop(self):
        self.stop_event.set()


def wait_until(pred, timeout=5.0):
    t_end = time.time() + timeout
    while time.time() < t_end:
        if pred():
            return True
        time.sleep(0.01)
    return False


def test_rules():
    st = StateStore()
    m = JobManager(st, max_active=2, job_factory=FakeJob)
    sub = lambda c, cat: m.submit(c, cat, [], "USA", 100, "serper", 12)  # noqa: E731

    a = sub("user-a", "Finance")
    b = sub("user-b", "RPO")
    c = sub("user-c", "Finance")        # same category as a: must wait for a
    d = sub("user-d", "Advisory")       # no free slot: queued behind c
    assert (a.status, b.status) == ("running", "running") or \
        wait_until(lambda: a.status == b.status == "running")
    assert c.status == "queued" and d.status == "queued"
    assert m.queue_position(c.id) == 1 and m.queue_position(d.id) == 2
    assert a.search_share() == 2 and a.crawl_workers == config.job_crawl_workers()
    # the first job takes what is free; the second what a leaves, at least
    # an equal slice
    assert b.crawl_workers == config.job_crawl_workers(a.crawl_workers)
    fair = max(8, min(config.CRAWL_WORKERS,
                      config.GLOBAL_CRAWL_WORKERS // config.MAX_ACTIVE_JOBS))
    assert a.crawl_workers >= b.crawl_workers >= fair
    print("  slots + FIFO + fair share: OK")

    # b finishes -> d (free category) overtakes c (its category is busy)
    b.release.set()
    assert wait_until(lambda: d.status == "running"), d.status
    assert c.status == "queued"
    print("  a blocked category never blocks other users' jobs: OK")

    # a finishes -> c finally starts
    a.release.set()
    assert wait_until(lambda: c.status == "running"), c.status
    print("  same-category job starts when the first one finishes: OK")

    # per-user limit, queue cap
    try:
        sub("user-c", "Law_Firms")
        raise AssertionError("second active job for one user was accepted")
    except JobRejected as exc:
        assert exc.status == 409
    old_cap = config.MAX_QUEUED_JOBS
    config.MAX_QUEUED_JOBS = 1
    try:
        e = sub("user-e", "Law_Firms")
        assert e.status == "queued"
        try:
            sub("user-f", "Forms")
            raise AssertionError("queue cap not enforced")
        except JobRejected as exc:
            assert exc.status == 503
    finally:
        config.MAX_QUEUED_JOBS = old_cap
    print("  per-user limit (409) + queue cap (503): OK")

    # stop while queued: finished immediately, never started
    stopped = m.stop("user-e")
    assert stopped is e and e.status == "stopped" and e.thread is None
    assert m.stop("user-x") is None                      # nothing of theirs
    assert m.stop("user-e", c.id) is None               # not their job
    print("  stop while queued / only your own job: OK")

    # stop while running -> slot freed, next queued job starts
    g = sub("user-g", "Forms")
    assert g.status == "queued"
    m.stop("user-c")
    assert wait_until(lambda: g.status == "running"), g.status
    d.release.set()
    g.release.set()
    assert wait_until(lambda: not m.active)
    print("  stop while running frees the slot: OK")

    # a job thread that dies without its finish callback releases its slot
    h = sub("user-h", "Other_B2B")
    assert wait_until(lambda: h.status == "running")
    h.on_finish = None
    h.release.set()
    assert wait_until(lambda: not m.active, 5.0), "dead job kept its slot"
    print("  reaper releases a slot held by a dead thread: OK")

    # bounded history
    old = config.JOB_HISTORY
    config.JOB_HISTORY = 5
    try:
        for k in range(12):
            j = sub(f"hist-{k}", "3D_Studios")
            j.release.set()
            assert wait_until(lambda j=j: j.status == "completed")
            assert wait_until(lambda j=j: j.id not in m.active)
        assert len(m.history) <= 5 and len(m.jobs) <= 5 + len(m.active) + len(m.queue)
    finally:
        config.JOB_HISTORY = old
    print("  finished-job history is bounded: OK")
    m.stop_all(timeout=2)


# --------------------------------------------------------------------------- #
# Part 2: real engine, concurrent jobs
# --------------------------------------------------------------------------- #

def test_real_concurrent_jobs():
    # Serper answers in ~0.2 s and the account allows 50 req/s, so the jobs
    # together WANT more than SERPER_CONCURRENCY requests in flight - the
    # account-wide cap must be what stops them.
    sim.TIME_SCALE = 0.15
    sim.ACCOUNT.rps = 50
    sim.install()
    config.ANALYZE_PROCESSES = 0      # parse in-thread: keeps the test light
    st = StateStore()
    m = JobManager(st, max_active=3)
    cats = ["Law_Firms", "RPO", "Advisory", "Finance"]
    targets = {c: 150 for c in cats}
    js = {c: m.submit(f"user-{c}", c, [], "USA", targets[c], "serper", 30,
                      geo={"country": "USA"}) for c in cats}
    assert sum(j.status == "queued" for j in js.values()) == 1
    t0 = time.time()
    assert wait_until(lambda: not m.active and not m.queue, 240), \
        {c: (j.status, j.category_count()) for c, j in js.items()}
    dt = time.time() - t0
    for c, j in js.items():
        assert j.status == "completed", (c, j.status, j.stop_reason)
        assert len(st.records[c]) == targets[c], (c, len(st.records[c]))
    # no business twice across the concurrent jobs
    doms, names = set(), set()
    for c in cats:
        for r in st.records[c]:
            d = norm_domain(r.get("Official Website") or "")
            n = norm_name(r.get("Company Name") or "")
            assert not d or d not in doms, f"duplicate domain across jobs: {d}"
            assert n not in names, f"duplicate name across jobs: {n}"
            if d:
                doms.add(d)
            names.add(n)
    print(f"  {len(cats)} jobs on 3 slots completed in {dt:.1f}s, "
          f"{sum(targets.values())} records, 0 cross-job duplicates: OK")
    # account-wide Serper concurrency held across all jobs - and was really
    # under pressure (otherwise this check would prove nothing)
    peak = sim.ACCOUNT.max_inflight
    assert peak <= config.SERPER_CONCURRENCY, peak
    assert peak >= config.SERPER_CONCURRENCY - 2, f"cap never exercised (peak {peak})"
    print(f"  Serper in flight (all jobs together) peaked at {peak}, cap "
          f"{config.SERPER_CONCURRENCY}; {sim.ACCOUNT.throttled} x 429: OK")
    # threads of finished jobs are gone (no leak across jobs)
    time.sleep(2.0)
    extra = [t.name for t in threading.enumerate()
             if t.name.startswith(("job-", "search-", "crawl"))]
    assert not extra, f"threads left behind by finished jobs: {extra[:8]}"
    print("  no threads left behind by finished jobs: OK")
    # checkpoint + Excel hold everything
    st.flush()
    reloaded = StateStore.load(config.STATE_PATH)
    for c in cats:
        assert len(reloaded.records[c]) == targets[c], c
    from openpyxl import load_workbook
    for c in cats:
        path = os.path.join(config.OUTPUT_DIR, st.all_categories()[c]["file"])
        wb = load_workbook(path, read_only=True)
        rows = sum(1 for row in wb.active.iter_rows(min_row=2, values_only=True)
                   if any(v not in (None, "") for v in row))
        wb.close()
        assert rows == targets[c], (c, rows)
    leftovers = [f for f in os.listdir(config.OUTPUT_DIR) if ".tmp" in f]
    assert not leftovers, leftovers
    print("  checkpoint + Excel files contain every record, no temp leftovers: OK")
    m.stop_all(timeout=2)


def main():
    print("Part 1: scheduling rules")
    test_rules()
    print("Part 2: real engine, concurrent jobs")
    test_real_concurrent_jobs()
    print("\nALL CONCURRENCY TESTS PASSED")
    print("outputs in:", _tmp)


if __name__ == "__main__":
    main()
