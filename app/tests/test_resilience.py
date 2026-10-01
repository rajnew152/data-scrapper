"""Fault-tolerance tests: ONE ERROR MUST NOT STOP THE WHOLE COLLECTION JOB.

Runs the REAL engine (planner, dedup registry, validation, checkpointing,
Excel export) against a simulated Serper + simulated business websites, with
deliberate failures injected at every level:

    search      random timeouts, a 12-request Serper outage, malformed result
                rows, duplicate-only queries
    crawl       failed URLs, parser exceptions, worker crashes, tasks stuck
                far beyond every timeout (watchdog must abandon them), a burst
                of stuck tasks that jams the pool (must be replaced)
    controller  an exception inside the controller loop (RECOVERING)
    process     a crash at ~455/2000 followed by a restart from the on-disk
                checkpoint

No Serper credits and no network are used. Run (from app/):
    python tests/test_resilience.py
"""

from __future__ import annotations

import hashlib
import os
import random
import shutil
import sys
import tempfile
import threading
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend import config  # noqa: E402

_tmp = tempfile.mkdtemp(prefix="resilience_test_")
config.OUTPUT_DIR = _tmp
config.STATE_PATH = os.path.join(_tmp, "demo_state.json")
config.ANALYZE_PROCESSES = 0       # in-thread parsing so faults can be injected
config.STATE_SAVE_INTERVAL = 0.5   # tight checkpoints: the crash test resumes ~455
config.STUCK_TASK_S = 3
config.WATCHDOG_STALL_S = 2
config.WATCHDOG_BOTTLENECK_S = 6
config.MAX_CREDITS_PER_RECORD = 0  # exhaustion must come from the search space

from backend.collector import analysis, discovery, engine  # noqa: E402
from backend.collector.engine import CollectionJob, StateStore  # noqa: E402
from backend.collector.normalize import norm_domain, norm_name, norm_phone  # noqa: E402
from backend.collector.search import (ProviderDisabled, SearchError,  # noqa: E402
                                      SearchProvider)

discovery.RETRY_BASE_S = 0.3       # failed searches re-queued after 0.3 s, 0.6 s
RELEASE = threading.Event()        # frees simulated stuck fetches at teardown


def h(*parts) -> int:
    return int(hashlib.md5("|".join(map(str, parts)).encode()).hexdigest()[:12], 16)


# --------------------------------------------------------------------------- #
# Simulated world: a finite universe of real-looking, unique businesses
# --------------------------------------------------------------------------- #

FIRST = ["Adler", "Brooks", "Carver", "Dalton", "Ellis", "Fisher", "Garner",
         "Hayes", "Irving", "Jensen", "Keller", "Lawson", "Mercer", "Nolan",
         "Ortiz", "Parker", "Quinn", "Ramsey", "Sutton", "Tanner", "Upton",
         "Vaughn", "Walker", "Xavier", "Yates", "Zimmer", "Harlow", "Kendall",
         "Monroe", "Prescott"]
SECOND = ["Talent", "Workforce", "Hire", "Career", "People", "Staff", "Placement",
          "Recruit", "Crew", "Team", "Skill", "Labor", "Pro", "Prime", "Bridge",
          "Summit", "Harbor", "Pioneer", "Keystone", "Beacon", "Compass",
          "Anchor", "Vertex", "Liberty", "Frontier", "Heritage", "Evergreen",
          "Northstar", "Blueline", "Redwood"]
KIND = ["Staffing", "Recruiters", "Staffing Partners", "Search Group",
        "Personnel", "Placements", "Talent Agency", "Employment Services",
        "Workforce Solutions", "Hiring Partners"]
CITIES = [("Austin", "Texas", "78701"), ("Dallas", "Texas", "75201"),
          ("Chicago", "Illinois", "60601"), ("Denver", "Colorado", "80202"),
          ("Seattle", "Washington", "98101"), ("Atlanta", "Georgia", "30303")]


class World:
    def __init__(self, size: int):
        self.size = size

    def biz(self, i: int) -> dict:
        i %= self.size
        name = (f"{FIRST[i % 30]} {SECOND[(i // 30) % 30]} "
                f"{KIND[(i // 900) % 10]}")
        dom = f"{FIRST[i % 30].lower()}{SECOND[(i // 30) % 30].lower()}{i}.example"
        city, state, z = CITIES[i % len(CITIES)]
        return {
            "id": i, "name": name, "domain": dom,
            "has_site": h("site", i) % 10 < 6,          # 60% have a website
            "phone": f"+1 {3000000000 + i}",
            "email": f"contact@{dom}",
            "address": f"{100 + i % 800} Main Street, {city}, {state} {z}, USA",
        }

    def ids_for(self, query: str, page: int, n: int = 10) -> list[int]:
        """Businesses a search returns: drawn from a window that depends on
        the query's location, so searches in one place overlap (realistic
        duplicates) and new places surface new businesses."""
        geo = query.split(" in ", 1)[1] if " in " in query else query
        base = h("geo", geo.lower()) % self.size
        rng = random.Random(h("q", query.lower(), page))
        window = min(self.size, 320)
        return [(base + k) % self.size for k in rng.sample(range(window), n)]


class SimProvider(SearchProvider):
    """Serper stand-in with injected failures."""
    name = "sim"
    supports_places = True

    def __init__(self, world: World, outage=(40, 52), fail_every=23,
                 bad_row_every=31, latency=(0.08, 0.25)):
        self.world = world
        self.outage = outage
        self.fail_every = fail_every
        self.bad_row_every = bad_row_every
        self.latency = latency
        self.calls = 0
        self.credits_used = 0
        self.ok_count = 0
        self.fail_count = 0
        self.issued: list[tuple[str, str, int]] = []
        self.lock = threading.Lock()

    @property
    def requests_made(self) -> int:
        return self.calls

    def _call(self, kind: str, query: str, page: int) -> list[int]:
        with self.lock:
            self.calls += 1
            n = self.calls
            self.issued.append((kind, query.lower(), page))
        time.sleep(random.uniform(*self.latency))
        if self.outage[0] <= n < self.outage[1]:
            with self.lock:
                self.fail_count += 1
            raise SearchError("simulated Serper outage (HTTP 503)")
        if self.fail_every and n % self.fail_every == 0:
            with self.lock:
                self.fail_count += 1
            raise SearchError("simulated Serper timeout")
        with self.lock:
            self.credits_used += 1
            self.ok_count += 1
        if "temp agency" in query.lower():       # duplicate-only query
            return list(range(10))
        return self.world.ids_for(query, page)

    def organic(self, query, max_results):
        return self.organic_page(query, 1)

    def organic_page(self, query, page):
        ids = self._call("organic", query, page)
        urls = [f"https://{self.world.biz(i)['domain']}/"
                for i in ids if self.world.biz(i)["has_site"]]
        return urls + ["https://www.linkedin.com/company/some-firm"]

    def places_page(self, query, page):
        ids = self._call("places", query, page)
        rows: list = []
        for i in ids:
            b = self.world.biz(i)
            row = {"title": b["name"], "phoneNumber": b["phone"],
                   "cid": str(900000 + b["id"]), "address": b["address"],
                   "category": "Employment agency"}
            if b["has_site"]:
                row["website"] = f"https://{b['domain']}"
            rows.append(row)
        if self.bad_row_every and self.calls % self.bad_row_every == 0:
            rows.insert(3, "<<malformed row>>")   # processing error mid-results
        return rows


class SimFetcher:
    """Website stand-in with injected crawl failures."""

    def __init__(self, world: World, stuck_at=(), jam_after=None, jam_count=0,
                 jam_secs=0.0):
        self.world = world
        self.stuck_at = set(stuck_at)     # fetch numbers that hang
        self.jam_after, self.jam_count, self.jam_secs = jam_after, jam_count, jam_secs
        self.fetches = 0
        self.lock = threading.Lock()

    def _id(self, url: str) -> int:
        dom = norm_domain(url)
        digits = "".join(ch for ch in dom.split(".")[0] if ch.isdigit())
        return int(digits) if digits else -1

    def fetch_raw(self, url: str):
        with self.lock:
            self.fetches += 1
            n = self.fetches
        i = self._id(url)
        time.sleep(random.uniform(0.05, 0.2))
        if i < 0:
            return None, url, "connection/DNS error"
        if n in self.stuck_at:                       # far beyond every timeout
            RELEASE.wait(config.STUCK_TASK_S * 4)
        if self.jam_after is not None and self.jam_after <= n < self.jam_after + self.jam_count:
            RELEASE.wait(self.jam_secs)              # burst that jams the pool
        if i % 53 == 0:
            raise RuntimeError(f"simulated worker crash on {url}")
        if i % 19 == 0:
            return None, url, "connection/DNS error"
        b = self.world.biz(i)
        html = f"""<html><head><title>{b['name']} | Staffing Agency</title></head>
          <body><h1>{b['name']}</h1>
          <a href="mailto:{b['email']}">Email us</a> <a href="tel:{b['phone']}">Call</a>
          <address>{b['address']}</address>
          <p>Recruitment and staffing services for growing companies.</p></body></html>"""
        return html.encode(), url, ""


_real_analyze_html = analysis.analyze_html


def _faulty_analyze_html(raw, url, *a, **kw):
    if url.split("//")[-1].split(".")[0].endswith("47"):   # ~1% of sites
        raise ValueError("simulated parser crash")
    return _real_analyze_html(raw, url, *a, **kw)


analysis.analyze_html = _faulty_analyze_html


# --------------------------------------------------------------------------- #
# Harness
# --------------------------------------------------------------------------- #

class Monitor(threading.Thread):
    """Samples progress + status: the timeline shown in the report."""

    def __init__(self, job: CollectionJob, every: float = 0.25):
        super().__init__(daemon=True)
        self.job, self.every = job, every
        self.samples: list[tuple[float, int, str]] = []
        self.stop_flag = threading.Event()

    def run(self):
        t0 = time.time()
        while not self.stop_flag.is_set():
            self.samples.append((time.time() - t0, self.job.category_count(),
                                 self.job.status))
            if self.job.status in engine.TERMINAL_STATUSES:
                break
            time.sleep(self.every)

    def timeline(self, marks=(0.1, 0.25, 0.5, 0.75, 0.9, 1.0)) -> str:
        out, last_status = [], None
        target = self.job.target
        pending = [int(m * target) for m in marks]
        for t, n, s in self.samples:
            if s != last_status and s in ("recovering", "waiting"):
                out.append(f"{n}/{target} [{s.upper()} {t:.1f}s]")
            elif pending and n >= pending[0]:
                while pending and n >= pending[0]:
                    pending.pop(0)
                out.append(f"{n}/{target}")
            last_status = s
        out.append(self.job.status.upper())
        return " -> ".join(out)


def make_job(state, world, target, provider, fetcher, category="Recruitment",
             location="USA", geo=None, max_queries=12):
    job = CollectionJob(state, category, [], location, target, "serper",
                        max_queries, geo=geo or {"country": location})
    job.fetcher = fetcher
    engine.make_provider = lambda *_: provider
    return job


def run(job, timeout=600, controller_fault_at=None) -> Monitor:
    if controller_fault_at is not None:          # one controller-loop exception
        real = job._want_organic
        calls = {"n": 0}

        def flaky():
            calls["n"] += 1
            if calls["n"] == controller_fault_at:
                raise RuntimeError("simulated controller bug")
            return real()
        job._want_organic = flaky
    mon = Monitor(job)
    mon.start()
    job.start()
    job.thread.join(timeout)
    mon.stop_flag.set()
    mon.join(2)
    assert not job.thread.is_alive(), "job did not finish (hang)"
    return mon


def check_records(state, category, expected):
    recs = state.records[category]
    assert len(recs) == expected, f"{len(recs)} records, expected {expected}"
    doms = [norm_domain(r["Official Website"]) for r in recs if r["Official Website"]]
    names = [norm_name(r["Company Name"]) for r in recs]
    phones = [norm_phone(r["Business Phone"]) for r in recs if r["Business Phone"]]
    assert len(set(doms)) == len(doms), "duplicate domains stored"
    assert len(set(names)) == len(names), "duplicate names stored"
    assert len(set(phones)) == len(phones), "duplicate phones stored"
    for r in recs:   # every record is a genuine simulated business, not invented
        assert r["Company Name"] and (r["Business Phone"] or r["Business Email"])


def excel_rows(state, category) -> int:
    from openpyxl import load_workbook
    path = os.path.join(config.OUTPUT_DIR, state.all_categories()[category]["file"])
    wb = load_workbook(path, read_only=True)
    try:
        return sum(1 for row in wb.active.iter_rows(min_row=2, values_only=True)
                   if any(v not in (None, "") for v in row))
    finally:
        wb.close()


def report(name, job, mon, t):
    c, p = job.counters, job.perf()
    print(f"\n=== {name}: {job.status.upper()} {job.category_count()}/{job.target} "
          f"in {t:.1f}s ===")
    print("  timeline:", mon.timeline())
    print(f"  faults survived: search failures {c['search_failures']} "
          f"(retried {c['search_retries']}, finally failed {c['searches_failed']}), "
          f"failed URLs {c['failed']}, worker errors {c['worker_errors']}, "
          f"abandoned stuck tasks {c['abandoned_tasks']}, pool restarts "
          f"{c['pool_restarts']}, controller errors {c['controller_errors']}, "
          f"recoveries {c['recoveries']}, strategy escalations "
          f"{c['strategy_escalations']}, watchdog actions {c['watchdog_actions']}")
    print(f"  metrics: {p['serper_requests_per_min']} search req/min, "
          f"avg search {p['serper_avg_ms']}ms, search success "
          f"{p['search_success_rate']}, {p['new_results_per_search']} new/search, "
          f"result dup rate {p['result_duplicate_rate']}, crawl avg "
          f"{p['crawl_avg_ms']}ms, validate avg {p['validation_avg_ms']}ms, "
          f"{p['valid_per_min']} valid/min, {p['valid_per_credit']} valid/credit, "
          f"failed ops {p['failed_operations']}, retries {p['retries']}")


# --------------------------------------------------------------------------- #
# Tests
# --------------------------------------------------------------------------- #

def test_serper_client_retries():
    """HTTP-level: transient failures retried with backoff, permanent ones not,
    auth/credit errors disable the provider."""
    import requests

    from backend.collector import search as search_mod
    from backend.collector.search import SerperProvider

    class Resp:
        def __init__(self, status, body=None, text="", headers=None):
            self.status_code, self._body = status, body
            self.text, self.headers = text, headers or {}

        def json(self):
            if isinstance(self._body, Exception):
                raise self._body
            return self._body

    class Session:
        def __init__(self, script):
            self.script = list(script)
            self.calls = 0

        def post(self, *a, **kw):
            self.calls += 1
            item = self.script.pop(0)
            if isinstance(item, Exception):
                raise item
            return item

    real_sleep = search_mod.time.sleep
    search_mod.time.sleep = lambda s: None
    config.SERPER_MAX_ATTEMPTS = 6
    try:
        p = SerperProvider("test-key")
        p.session = Session([requests.Timeout(), requests.ConnectionError("DNS"),
                             Resp(503, text="unavailable"),
                             Resp(200, ValueError("bad json")),
                             Resp(429, headers={"Retry-After": "2"}),
                             Resp(200, {"organic": [{"link": "https://a.example"}]})])
        assert p.organic("q", 10) == ["https://a.example"]
        assert p.session.calls == 6 and p.retry_count == 5, (p.session.calls, p.retry_count)

        p = SerperProvider("test-key")
        p.session = Session([Resp(400, text="bad query")])
        try:
            p.organic("q", 10)
            raise AssertionError("permanent 400 must not be retried")
        except SearchError:
            assert p.session.calls == 1

        p = SerperProvider("test-key")
        p.session = Session([Resp(403, text="Not enough credits")])
        try:
            p.organic("q", 10)
            raise AssertionError("403 must disable the provider")
        except ProviderDisabled:
            pass
        p = SerperProvider("test-key")
        p.session = Session([Resp(500, text="error: credit check failed")] * 6)
        try:
            p.organic("q", 10)
        except SearchError:
            pass   # a 5xx mentioning "credit" is transient, not a disable
        assert not p.disabled_reason
    finally:
        search_mod.time.sleep = real_sleep
        config.SERPER_MAX_ATTEMPTS = 4
    print("serper client retry/backoff tests passed")


def run_target(target: int, world_size: int = 9000, **fetch_kw) -> CollectionJob:
    state = StateStore()
    world = World(world_size)
    job = make_job(state, world, target, SimProvider(world, outage=(10, 22)),
                   SimFetcher(world, stuck_at=(25, 60, 95), **fetch_kw))
    t0 = time.time()
    mon = run(job, controller_fault_at=15)
    report(f"TARGET {target}", job, mon, time.time() - t0)
    assert job.status == "completed", (job.status, job.stop_reason, job.error)
    check_records(state, "Recruitment", target)
    assert excel_rows(state, "Recruitment") == target
    c = job.counters
    assert c["search_failures"] >= 10, "outage/timeouts not exercised"
    if target >= 500:   # enough crawls that injected crawl faults always land
        assert c["worker_errors"] >= 1 and c["failed"] >= 1
    assert c["controller_errors"] == 1 and c["recoveries"] >= 1
    return job


def test_pool_jam():
    """A burst of tasks stuck far beyond every timeout jams a small pool: the
    watchdog abandons them, replaces the pool, and the run completes."""
    saved = config.CRAWL_WORKERS
    config.CRAWL_WORKERS = 8
    try:
        job = run_target(500, jam_after=120, jam_count=6, jam_secs=30)
    finally:
        config.CRAWL_WORKERS = saved
    assert job.counters["abandoned_tasks"] >= 4, job.counters["abandoned_tasks"]
    assert job.counters["pool_restarts"] >= 1, "jammed pool was not replaced"


def test_crash_and_resume():
    """Crash at ~455/2000, restart from the on-disk checkpoint, finish at
    2000/2000 without re-running searches or losing records."""
    world = World(9000)
    image = os.path.join(_tmp, "crash_image.json")
    if os.path.exists(config.STATE_PATH):
        os.remove(config.STATE_PATH)

    state_a = StateStore()
    job_a = make_job(state_a, world, 2000, SimProvider(world), SimFetcher(world))
    saved_conc = config.SERPER_CONCURRENCY
    config.SERPER_CONCURRENCY = 3     # slower phase A: the crash lands near 455
    job_a.start()
    while True:   # "crash" once the on-disk checkpoint holds >= 455 records
        time.sleep(0.05)
        try:
            n_disk = len(StateStore.load(config.STATE_PATH).records["Recruitment"])
        except Exception:
            n_disk = 0
        if n_disk >= 455:
            shutil.copyfile(config.STATE_PATH, image)
            break
        assert job_a.thread.is_alive(), job_a.status
    job_a.stop()                      # the process "dies": its threads go away
    job_a.thread.join(30)
    config.SERPER_CONCURRENCY = saved_conc

    state_b = StateStore.load(image)  # what a restarted server loads
    resumed_from = len(state_b.records["Recruitment"])
    before = {norm_domain(r["Official Website"]) or r["Company Name"]
              for r in state_b.records["Recruitment"]}
    used_before = set(state_b.discovery["Recruitment"]["used"])
    provider_b = SimProvider(world, outage=(30, 40))
    job_b = make_job(state_b, world, 2000, provider_b, SimFetcher(world))
    issued_keys: list[str] = []
    real_next = discovery.DiscoveryPlanner.next_requests

    def spy(self, *a, **kw):   # every request the resumed planner issues
        reqs = real_next(self, *a, **kw)
        for r in reqs:
            if not r.meta.get("attempt") and not r.meta.get("resumed"):
                # a search = (cell, type, page): Places and organic of one
                # cell are different searches
                issued_keys.append(f"{r.phrase}@{r.geo}#"
                                   f"{'p' if r.kind == 'places' else 'o'}{r.page}")
        return reqs

    resume_lines: list[str] = []
    real_say = job_b.say

    def say(msg):
        if msg.startswith("Resuming:"):
            resume_lines.append(msg)
        real_say(msg)
    job_b.say = say

    t0 = time.time()
    discovery.DiscoveryPlanner.next_requests = spy
    try:
        mon = run(job_b)
    finally:
        discovery.DiscoveryPlanner.next_requests = real_next
    report(f"CRASH AT {resumed_from}/2000 -> RESTART -> RESUME", job_b, mon, time.time() - t0)
    assert resume_lines, "resume not announced"
    print("  resume log:", resume_lines[0][:130])
    assert job_b.status == "completed", (job_b.status, job_b.stop_reason, job_b.error)
    assert 455 <= resumed_from < 700, resumed_from
    check_records(state_b, "Recruitment", 2000)
    after = {norm_domain(r["Official Website"]) or r["Company Name"]
             for r in state_b.records["Recruitment"]}
    assert before <= after, "records from before the crash were lost"
    repeated = [k for k in issued_keys if k in used_before]
    assert not repeated, f"resumed run repeated {len(repeated)} searches: {repeated[:3]}"
    assert excel_rows(state_b, "Recruitment") == 2000
    reloaded = StateStore.load(config.STATE_PATH)
    assert len(reloaded.records["Recruitment"]) == 2000
    meta = reloaded.jobs["Recruitment"]
    assert meta["status"] == "completed" and meta["target"] == 2000, meta
    print(f"  resume: {resumed_from} records restored, {len(issued_keys)} new "
          f"searches issued, 0 repeated; checkpoint job meta: {meta['status']} "
          f"{meta['collected']}/{meta['target']}")


def test_exhaustion_is_honest():
    """A market with only ~150 reachable businesses and a target of 500:
    the job walks every strategy tier, then ends EXHAUSTED with an exact
    reason - no fake records, no infinite loop."""
    state = StateStore()
    world = World(150)
    job = make_job(state, world, 500, SimProvider(world, outage=(0, 0)),
                   SimFetcher(world), location="Austin, Texas, USA",
                   geo={"city": "Austin", "state": "Texas", "country": "USA"})
    t0 = time.time()
    mon = run(job)
    report("EXHAUSTION (150 real businesses, target 500)", job, mon, time.time() - t0)
    n = len(state.records["Recruitment"])
    assert job.status == "exhausted", job.status
    assert n <= 150 and n > 50, n
    assert job.planner.tier_name == "deep" and job.counters["strategy_escalations"] >= 3
    assert "search strategy tier" in job.stop_reason
    check_records(state, "Recruitment", n)
    print(f"  honest stop: {n} genuine records (universe 150), reason: {job.stop_reason[:90]}")


def test_credits_run_out_then_topped_up():
    """The screenshot case: Serper answers 400 "Not enough credits" mid-run.
    The job must PAUSE (status waiting, progress saved) - not fail - and
    continue by itself once the balance is back, finishing at 100%."""
    class CreditLimited(SimProvider):
        def __init__(self, world, allowance):
            super().__init__(world, outage=(0, 0), fail_every=0, bad_row_every=0)
            self.allowance = allowance
            self.topped_up = threading.Event()

        def _call(self, kind, query, page):
            with self.lock:
                empty = self.credits_used >= self.allowance and not self.topped_up.is_set()
            if empty:
                raise ProviderDisabled('Serper HTTP 400: {"message":"Not enough credits",'
                                       '"statusCode":400}', "credits")
            return super()._call(kind, query, page)

        def balance(self):
            if self.topped_up.is_set():
                return 500
            with self.lock:
                return max(0, self.allowance - self.credits_used)

    saved = engine.CREDIT_POLL_S
    engine.CREDIT_POLL_S = 0.5
    try:
        state = StateStore()
        world = World(9000)
        provider = CreditLimited(world, allowance=8)   # runs out under any search mix
        job = make_job(state, world, 300, provider, SimFetcher(world))
        mon = Monitor(job)
        mon.start()
        job.start()
        t_end = time.time() + 120
        while job.status != "waiting" and time.time() < t_end:
            assert job.status in engine.ACTIVE_STATUSES, (job.status, job.error)
            time.sleep(0.05)
        assert job.status == "waiting", job.status
        paused_at = job.category_count()
        time.sleep(2.0)                       # several balance checks at 0
        assert job.status == "waiting" and job.thread.is_alive(), job.status
        assert os.path.exists(config.STATE_PATH)
        on_disk = len(StateStore.load(config.STATE_PATH).records["Recruitment"])
        assert on_disk >= paused_at, (on_disk, paused_at)   # progress saved
        provider.topped_up.set()              # user adds credits
        job.thread.join(120)
        mon.stop_flag.set()
        mon.join(2)
    finally:
        engine.CREDIT_POLL_S = saved
    report("CREDITS RUN OUT -> PAUSED -> TOP-UP -> RESUMED", job, mon, 0.0)
    lines = list(job.log_lines)
    assert job.status == "completed", (job.status, job.error)
    assert paused_at < 300
    check_records(state, "Recruitment", 300)
    print(f"  paused at {paused_at}/300 (saved {on_disk} on disk), resumed after "
          f"top-up, completed 300/300; "
          f"{sum('Still waiting' in x for x in lines)} balance checks while paused")


def test_provider_disabled_fails_exactly():
    class Dead(SimProvider):
        def places_page(self, query, page):
            raise ProviderDisabled("Serper HTTP 401: Unauthorized - invalid API key")
        organic_page = places_page

    state = StateStore()
    world = World(500)
    job = make_job(state, world, 50, Dead(world), SimFetcher(world))
    run(job)
    assert job.status == "failed" and "invalid API key" in job.error, (job.status, job.error)
    print(f"\nprovider disabled -> FAILED with exact reason: {job.error}")


def main():
    random.seed(7)
    try:
        test_serper_client_retries()
        run_target(100)
        test_pool_jam()
        run_target(2000)
        test_crash_and_resume()
        test_exhaustion_is_honest()
        test_credits_run_out_then_topped_up()
        test_provider_disabled_fails_exactly()
    finally:
        RELEASE.set()
    print("\nALL RESILIENCE TESTS PASSED")
    print("outputs in:", config.OUTPUT_DIR)


if __name__ == "__main__":
    main()
